from .base import DocumentLoader, LoadedDocument, DocumentPage
from typing import List
import structlog

logger = structlog.get_logger()


class DOCXLoader(DocumentLoader):
    """Load DOCX files using python-docx."""

    def supported_extensions(self) -> List[str]:
        return ["docx"]

    def load(self, file_path: str) -> LoadedDocument:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required: pip install python-docx")

        logger.info("Loading DOCX", path=file_path)
        doc = Document(file_path)

        # Group paragraphs into logical pages (~50 paragraphs each)
        all_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        page_size = 50
        pages: List[DocumentPage] = []

        for i in range(0, len(all_paragraphs), page_size):
            chunk = all_paragraphs[i: i + page_size]
            page_num = i // page_size + 1
            pages.append(DocumentPage(
                content="\n\n".join(chunk),
                page_number=page_num,
                metadata={"page": page_num},
            ))

        raw_meta = {}
        if doc.core_properties:
            raw_meta = {
                "title": doc.core_properties.title or "",
                "author": doc.core_properties.author or "",
                "description": doc.core_properties.description or "",
            }

        logger.info("DOCX loaded", pages=len(pages), paragraphs=len(all_paragraphs))
        return LoadedDocument(
            pages=pages,
            total_pages=len(pages),
            format="docx",
            raw_metadata=raw_meta,
        )
